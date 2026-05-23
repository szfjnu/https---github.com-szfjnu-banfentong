from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 创建文档
doc = Document()

# 设置标题
title = doc.add_heading('📖 班分通小程序使用指南（零基础操作手册）', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.size = Pt(16)
title.runs[0].font.bold = True

# 添加目录
doc.add_heading('📋 目录', level=1)
doc.add_paragraph('1. 【基础入门】如何进入小程序 & 加入班级')
doc.add_paragraph('2. 【班主任视图】完整功能操作指南')
doc.add_paragraph('3. 【科任老师视图】功能操作指南')
doc.add_paragraph('4. 【学生班委视图】功能操作指南')
doc.add_paragraph('5. 【学生视图】功能操作指南')
doc.add_paragraph('6. 【家长视图】功能操作指南')
doc.add_paragraph('7. 【常见问题解答】')
doc.add_paragraph('8. 【注意事项】')

# 添加分页
doc.add_page_break()

# ========== 一、基础入门 ==========
doc.add_heading('一、基础入门：三步搞定', level=1)

doc.add_heading('📱 第一步：进入小程序', level=2)
doc.add_paragraph('1. 打开微信，点击底部「发现」')
doc.add_paragraph('2. 点击「小程序」')
doc.add_paragraph('3. 在搜索框输入「班分通」，点击搜索')
doc.add_paragraph('4. 点击「班分通」小程序图标进入')

doc.add_heading('📱 第二步：选择角色', level=2)
doc.add_paragraph('1. 进入小程序后，点击「开始使用」')
doc.add_paragraph('2. 选择您的角色：班主任 / 科任老师 / 学生 / 家长')
doc.add_paragraph('3. 点击「确认」')

doc.add_heading('📱 第三步：加入班级', level=2)
doc.add_paragraph('1. 点击「加入班级」')
doc.add_paragraph('2. 输入班主任提供的班级码（例如：BFT2024）')
doc.add_paragraph('3. 点击「确定」，即可加入班级')

# 添加分页
doc.add_page_break()

# ========== 二、班主任视图 ==========
doc.add_heading('二、班主任视图', level=1)
doc.add_paragraph('🎯 您的核心权限：管理整个班级')

doc.add_heading('🏠 首页功能', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '操作路径'
hdr_cells[2].text = '效果说明'
rows = [
    ('班级积分概览', '首页 → 顶部卡片', '查看班级总积分、平均分、排名第一的学生'),
    ('今日待办', '首页 → 待办列表', '查看需要处理的审核任务'),
    ('快捷入口', '首页 → 底部图标', '快速进入积分管理、宿舍管理、通知发布'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]
    row_cells[2].text = row[2]

doc.add_heading('📊 积分管理', level=2)
doc.add_paragraph('操作路径：首页 → 积分管理')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('积分规则设置', '设置基础分（默认100分）、加减分项目'),
    ('积分记录', '查看所有学生的积分变动记录'),
    ('批量导入', '从Excel导入学生信息和初始积分'),
    ('导出成绩单', '一键导出期末成绩单（Excel格式）'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_paragraph('如何修改积分？')
doc.add_paragraph('1. 点击「积分记录」')
doc.add_paragraph('2. 找到需要修改的学生')
doc.add_paragraph('3. 点击「修改积分」')
doc.add_paragraph('4. 输入分数和原因')
doc.add_paragraph('5. 点击「确定」')

doc.add_heading('👥 成员管理', level=2)
doc.add_paragraph('操作路径：首页 → 成员管理')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('添加成员', '手动添加学生或老师'),
    ('设置班委', '指定学生担任班长、学习委员等职务'),
    ('权限设置', '给班委分配管理权限（如：积分登记、考勤管理）'),
    ('移除成员', '将成员移出班级'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('🏠 宿舍管理', level=2)
doc.add_paragraph('操作路径：首页 → 宿舍管理')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('创建宿舍', '设置宿舍名称、入住学生'),
    ('卫生检查', '查看宿舍卫生评分记录'),
    ('设置预警线', '低于多少分需要提醒'),
    ('宿舍排行榜', '查看所有宿舍的积分排名'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('📢 通知发布', level=2)
doc.add_paragraph('操作路径：首页 → 通知发布')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('发布通知', '编辑通知内容，选择接收人群'),
    ('通知列表', '查看所有已发布的通知'),
    ('撤回通知', '删除已发布的通知'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('👤 我的', level=2)
doc.add_paragraph('操作路径：首页 → 我的')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('班级信息', '修改班级名称、年级等信息'),
    ('班级码', '查看或重置班级码'),
    ('消息设置', '开启/关闭消息提醒'),
    ('退出班级', '离开当前班级（谨慎操作）'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

# 添加分页
doc.add_page_break()

# ========== 三、科任老师视图 ==========
doc.add_heading('三、科任老师视图', level=1)
doc.add_paragraph('🎯 您的核心权限：管理学科积分')

doc.add_heading('🏠 首页功能', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '操作路径'
hdr_cells[2].text = '效果说明'
rows = [
    ('学科积分', '首页 → 学科积分卡片', '查看所教班级的学科积分情况'),
    ('待审核', '首页 → 待审核列表', '查看需要审核的积分申请'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]
    row_cells[2].text = row[2]

doc.add_heading('📊 学科积分管理', level=2)
doc.add_paragraph('操作路径：首页 → 学科积分')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('添加积分', '为学生添加学科相关积分（如：作业优秀+2分）'),
    ('积分记录', '查看所有学科积分变动'),
    ('学科统计', '按学科查看学生积分排名'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_paragraph('如何给学生加学科积分？')
doc.add_paragraph('1. 点击「添加积分」')
doc.add_paragraph('2. 选择班级和学生')
doc.add_paragraph('3. 输入积分和原因（如：数学作业满分）')
doc.add_paragraph('4. 点击「提交」')

doc.add_heading('📝 学生评价', level=2)
doc.add_paragraph('操作路径：首页 → 学生评价')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('写评语', '对学生进行阶段性评价'),
    ('评价历史', '查看所有已写的评语'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

# 添加分页
doc.add_page_break()

# ========== 四、学生班委视图 ==========
doc.add_heading('四、学生班委视图', level=1)
doc.add_paragraph('🎯 您的核心权限：协助班主任管理班级')

doc.add_heading('🏠 首页功能', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '操作路径'
hdr_cells[2].text = '效果说明'
rows = [
    ('我的任务', '首页 → 任务卡片', '查看班主任分配的管理任务'),
    ('快捷入口', '首页 → 底部图标', '根据权限显示可操作的功能'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]
    row_cells[2].text = row[2]

doc.add_heading('📊 积分登记', level=2)
doc.add_paragraph('操作路径：首页 → 积分登记（需班主任授权）')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('登记积分', '为同学添加或扣除积分'),
    ('审核记录', '查看自己提交的积分申请状态'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_paragraph('如何登记积分？')
doc.add_paragraph('1. 点击「登记积分」')
doc.add_paragraph('2. 选择同学姓名')
doc.add_paragraph('3. 选择类型（加分/扣分）')
doc.add_paragraph('4. 输入分数和原因')
doc.add_paragraph('5. 点击「提交」（需班主任审核）')

doc.add_heading('📝 考勤管理', level=2)
doc.add_paragraph('操作路径：首页 → 考勤管理（需班主任授权）')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('今日考勤', '记录同学的出勤情况'),
    ('考勤统计', '查看本周考勤汇总'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('🧹 值日安排', level=2)
doc.add_paragraph('操作路径：首页 → 值日安排（需班主任授权）')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('本周值日表', '查看本周值日安排'),
    ('完成登记', '标记值日完成情况'),
    ('申请调换', '申请与其他同学调换值日'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

# 添加分页
doc.add_page_break()

# ========== 五、学生视图 ==========
doc.add_heading('五、学生视图', level=1)
doc.add_paragraph('🎯 您的核心权限：查看个人信息 & 参与班级活动')

doc.add_heading('🏠 首页功能', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '操作路径'
hdr_cells[2].text = '效果说明'
rows = [
    ('我的积分', '首页 → 积分卡片', '查看当前积分和排名'),
    ('班级通知', '首页 → 通知列表', '查看班主任发布的通知'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]
    row_cells[2].text = row[2]

doc.add_heading('📊 积分详情', level=2)
doc.add_paragraph('操作路径：首页 → 我的积分')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('积分变动记录', '查看所有加分/扣分记录'),
    ('班级排行榜', '查看全班积分排名'),
    ('我的荣誉', '查看获得的电子证书'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('🌳 树洞心声', level=2)
doc.add_paragraph('操作路径：首页 → 树洞')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('发表心声', '匿名发布心情或想法'),
    ('查看帖子', '浏览其他同学的心声'),
    ('点赞', '给认同的帖子点赞'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('🤝 志愿服务', level=2)
doc.add_paragraph('操作路径：首页 → 志愿服务')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('报名活动', '报名参加志愿服务活动'),
    ('我的服务', '查看已报名的活动'),
    ('服务时长', '查看累计志愿服务时长'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('👤 我的', level=2)
doc.add_paragraph('操作路径：首页 → 我的')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('个人信息', '查看/修改个人资料'),
    ('消息中心', '查看收到的消息'),
    ('设置', '开启/关闭通知提醒'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

# 添加分页
doc.add_page_break()

# ========== 六、家长视图 ==========
doc.add_heading('六、家长视图', level=1)
doc.add_paragraph('🎯 您的核心权限：关注孩子的在校表现')

doc.add_heading('🏠 首页功能', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '操作路径'
hdr_cells[2].text = '效果说明'
rows = [
    ('孩子积分', '首页 → 积分卡片', '查看孩子当前积分'),
    ('积分变动提醒', '首页 → 提醒列表', '查看积分变动通知'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]
    row_cells[2].text = row[2]

doc.add_heading('📊 孩子积分详情', level=2)
doc.add_paragraph('操作路径：首页 → 孩子积分')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('积分记录', '查看孩子所有积分变动原因'),
    ('积分趋势', '查看近一个月的积分变化'),
    ('班级排名', '查看孩子在班级的积分排名'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('🏠 宿舍情况', level=2)
doc.add_paragraph('操作路径：首页 → 宿舍情况（如有住宿生）')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('宿舍评分', '查看孩子宿舍的卫生评分'),
    ('宿舍成员', '查看宿舍其他成员'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('💬 家校沟通', level=2)
doc.add_paragraph('操作路径：首页 → 消息')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('班级通知', '查看班主任发布的通知'),
    ('私信老师', '给班主任发消息'),
    ('家长群', '进入班级家长群'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

doc.add_heading('👤 我的', level=2)
doc.add_paragraph('操作路径：首页 → 我的')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '按钮名称'
hdr_cells[1].text = '效果说明'
rows = [
    ('绑定孩子', '绑定多个孩子（如有）'),
    ('消息设置', '开启/关闭积分变动提醒'),
    ('常见问题', '查看使用帮助'),
]
for row in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = row[0]
    row_cells[1].text = row[1]

# 添加分页
doc.add_page_break()

# ========== 七、常见问题解答 ==========
doc.add_heading('七、常见问题解答', level=1)

questions = [
    ('Q1：忘记班级码怎么办？', '联系班主任获取班级码，或在微信群查看群公告。'),
    ('Q2：绑定孩子失败怎么办？', '请确认输入的班级码正确，或联系班主任核实孩子信息是否已录入系统。'),
    ('Q3：积分多久更新一次？', '实时更新！当有积分变动时，您会收到通知提醒。'),
    ('Q4：可以同时加入多个班级吗？', '可以的！点击首页右上角「切换班级」即可。'),
    ('Q5：如何修改个人信息？', '进入「我的」页面，点击「个人信息」即可修改。'),
    ('Q6：积分可以兑换奖品吗？', '目前体验版暂未开通兑换功能，正式版将增加此功能。'),
    ('Q7：数据安全吗？', '非常安全！所有数据都存储在云端，并且严格保密。'),
    ('Q8：忘记操作怎么办？', '本指南涵盖所有操作，随时查看即可。'),
]

for q, a in questions:
    doc.add_paragraph(f'❓ {q}')
    doc.add_paragraph(f'**A**：{a}')

# 添加分页
doc.add_page_break()

# ========== 八、注意事项 ==========
doc.add_heading('八、注意事项', level=1)
doc.add_paragraph('⚠️ **重要提醒**：')
doc.add_paragraph('1. 请妥善保管班级码，不要随意分享给非班级成员')
doc.add_paragraph('2. 积分变动需要班主任审核，请耐心等待')
doc.add_paragraph('3. 树洞功能请文明发言，禁止发布违规内容')
doc.add_paragraph('4. 如果遇到问题，请先查看常见问题，或联系班主任')
doc.add_paragraph('5. 体验版功能可能会调整，正式版将更加完善')

doc.add_paragraph('🎉 **恭喜您！** 现在您已经掌握了班分通的全部使用方法。如果在使用过程中有任何问题，随时联系班主任或在群里提问！')

doc.add_paragraph('📝 **版本说明**')
doc.add_paragraph('当前版本：体验版 v1.0')
doc.add_paragraph('更新时间：2025年6月')
doc.add_paragraph('正式版即将上线，敬请期待！')

# 保存文档
doc.save('班分通使用指南.docx')
print('Word文档已生成成功！')
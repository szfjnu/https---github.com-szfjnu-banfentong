from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 创建文档
doc = Document()

# 设置中文字体
def set_chinese_font(run, font_name='宋体', font_size=10.5, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.bold = bold

# 标题
title = doc.add_heading('', level=0)
title_run = title.add_run('班分通小程序详细操作手册')
set_chinese_font(title_run, font_name='黑体', font_size=22, bold=True)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 版本信息
version = doc.add_paragraph()
version_run = version.add_run('版本：体验版 v1.0\n适用对象：班主任、科任老师、学生班委、学生、家长\n文档性质：软件可交付标准用户操作说明文档')
set_chinese_font(version_run, font_name='宋体', font_size=10.5)
version.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# 目录
toc = doc.add_heading('', level=1)
toc_run = toc.add_run('目录')
set_chinese_font(toc_run, font_name='黑体', font_size=16, bold=True)

toc_items = [
    '一、系统概述',
    '二、基础操作指南',
    '三、首页功能详解',
    '四、班级模块详解',
    '五、积分模块详解',
    '六、发现模块详解',
    '七、个人中心详解',
    '八、角色专属功能',
    '九、数据流转与业务逻辑',
    '十、常见问题处理',
    '十一、界面元素速查表'
]

for item in toc_items:
    p = doc.add_paragraph(item, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=12)

doc.add_page_break()

# ========== 一、系统概述 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('一、系统概述'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('1.1 产品定位'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('班分通是一款面向中小学班级管理的微信小程序，基于微信云开发架构，支持多角色协同管理。系统涵盖学生管理、积分评价、考勤记录、宿舍管理、值日安排、处分管理、志愿服务、技能证书、成长档案、社交互动等核心功能模块。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('1.2 角色体系'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('系统定义六种用户角色，各角色拥有不同的功能权限：'), font_name='宋体', font_size=10.5)

# 角色表格
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
headers = ['角色', '角色标识', '核心权限范围']
for i, header in enumerate(headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

roles = [
    ('管理员', 'admin', '系统全局管理、用户管理、数据迁移'),
    ('班主任', 'head_teacher', '班级全面管理、审批审核、规则配置'),
    ('科任老师', 'subject_teacher', '学科积分管理、学生评价、考勤记录'),
    ('班干部', 'class_cadre', '积分登记、考勤管理、值日检查（需授权）'),
    ('学生', 'student', '查看个人信息、参与活动、申请审批'),
    ('家长', 'parent', '查看孩子信息、接收通知、家校沟通')
]

for role in roles:
    row_cells = table.add_row().cells
    for i, text in enumerate(role):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('1.3 系统架构'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('小程序端采用分包架构，包含以下功能包：'), font_name='宋体', font_size=10.5)

architecture = [
    '主包（pages/）：首页、班级、积分、发现、我的五个主入口页面',
    '核心功能包（subPkg1/）：学生管理、班级管理、考勤管理、积分管理、志愿服务、值日管理、处分管理、审批中心、学期管理、技能证书',
    '宿舍管理包（subPkg2/）：宿舍楼栋、房间床位、卫生评分、宿舍统计、检查记录、预警配置',
    '社交互动包（subPkg3/）：校园闲鱼、心灵树洞、英雄台、微聊陪伴、聚光点、成绩管理',
    '教务管理包（subPkg4/）：权限授权、课表管理、座位管理、分组管理、管理员功能',
    '用户系统包（subPkg5/）：登录注册、加入班级、成长档案、预警配置、消息通知、设置中心'
]

for item in architecture:
    p = doc.add_paragraph(item, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('1.4 数据集合'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('系统核心数据集合包括：users（用户基础信息）、students（学生档案信息）、classes（班级信息）、user_class_relation（用户班级关联关系）、score_records（积分变动记录）、score_items（积分项目配置）、score_rules（积分规则版本）、dorms（宿舍信息）、dorm_score_accounts（宿舍积分账户）、dorm_inspection_records（宿舍检查记录）、attendance_records（考勤记录）、duty_schedules（值日安排）、duty_check_records（值日检查记录）、discipline_records（处分记录）、volunteer_records（志愿服务记录）、skill_certs（技能证书记录）、treehole_posts（树洞帖子）、hero_honors（英雄台荣誉）、flea_items（跳蚤市场物品）、semesters（学期配置）、notifications（系统通知）等。'), font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 二、基础操作指南 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('二、基础操作指南'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('2.1 首次使用流程'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('步骤1：进入小程序'), font_name='黑体', font_size=12, bold=True)

steps = [
    '打开微信应用',
    '点击底部导航栏「发现」',
    '点击「小程序」选项',
    '在搜索框输入「班分通」',
    '点击搜索结果中的「班分通」图标进入'
]

for i, step in enumerate(steps, 1):
    p = doc.add_paragraph()
    set_chinese_font(p.add_run(f'{i}. {step}'), font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('界面元素说明：搜索框位于小程序列表顶部，支持模糊搜索；小程序图标为班分通品牌标识，点击进入；如曾使用过，可在最近使用列表直接点击。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('步骤2：登录授权'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('首次进入将显示登录页面，点击「微信一键登录」按钮，授权获取微信昵称和头像，系统自动创建用户档案。'), font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('按钮说明：「微信一键登录」调用微信登录接口，获取openid完成认证；点击效果为显示加载状态，成功后跳转角色选择页面。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('步骤3：选择角色'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('在角色选择页面，点击对应角色卡片（班主任/科任老师/学生/家长），点击「确认」按钮。'), font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('界面元素：角色卡片显示角色图标和名称，点击选中后高亮显示；「确认」按钮为蓝色主按钮，点击后进入班级加入流程。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('步骤4：加入班级'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('点击「加入班级」按钮，输入班主任提供的班级码（6位字母数字组合），点击「确定」，系统验证班级码有效性，验证通过后进入小程序首页。'), font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('输入框说明：班级码输入框限制6位字符，自动转大写；「确定」按钮验证中显示加载状态，成功跳转首页；班级码无效时显示红色提示文字。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('2.2 底部导航栏'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('小程序底部固定显示五个主入口，点击切换：'), font_name='宋体', font_size=10.5)

nav_table = doc.add_table(rows=1, cols=4)
nav_table.style = 'Table Grid'
hdr_cells = nav_table.rows[0].cells
nav_headers = ['图标', '文字', '页面路径', '功能说明']
for i, header in enumerate(nav_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

nav_items = [
    ('🏠', '首页', 'pages/index/index', '查看班级概览、快捷操作'),
    ('👥', '班级', 'pages/student/student', '学生列表、班级功能导航'),
    ('📊', '积分', 'pages/score/score', '积分排名、积分记录'),
    ('🔍', '发现', 'pages/discover/discover', '成长探索、校园生活'),
    ('👤', '我的', 'pages/usercenter/usercenter', '个人中心、功能菜单')
]

for item in nav_items:
    row_cells = nav_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('交互说明：点击未选中标签切换页面，图标变为高亮状态；点击已选中标签页面刷新回到顶部；选中状态文字和图标变为紫色（#667eea）；未选中状态文字和图标为灰色（#999999）。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('2.3 通用界面元素'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('顶部导航栏'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('背景色为蓝色（#1890ff），标题文字为白色显示当前页面名称，左侧为返回按钮（箭头图标，点击返回上一页），右侧为更多按钮（三个点，显示页面选项菜单）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('卡片组件'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('圆角矩形白色背景，轻微阴影效果，包含标题、内容、操作按钮，点击可进入详情页面。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('列表组件'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('左侧图标/头像，中间标题和描述，右侧箭头/状态标签，点击展开或跳转。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('按钮类型'), font_name='黑体', font_size=12, bold=True)

btn_table = doc.add_table(rows=1, cols=3)
btn_table.style = 'Table Grid'
hdr_cells = btn_table.rows[0].cells
btn_headers = ['类型', '样式', '使用场景']
for i, header in enumerate(btn_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

btn_types = [
    ('主按钮', '蓝色背景，白色文字', '确认、提交、保存'),
    ('次按钮', '白色背景，蓝色边框', '取消、返回、次要操作'),
    ('危险按钮', '红色背景，白色文字', '删除、移除、警告操作'),
    ('文字按钮', '纯文字，蓝色', '链接、查看更多'),
    ('禁用按钮', '灰色背景', '不可点击状态')
]

for btn in btn_types:
    row_cells = btn_table.add_row().cells
    for i, text in enumerate(btn):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('表单元素'), font_name='黑体', font_size=12, bold=True)

form_elements = [
    '输入框：底部边框，聚焦时变蓝',
    '选择器：点击弹出底部选择面板',
    '开关：左右滑动切换',
    '单选/多选：圆形/方形选择框',
    '日期选择：点击弹出日期面板'
]

for element in form_elements:
    p = doc.add_paragraph(element, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 三、首页功能详解 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('三、首页功能详解'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.1 页面结构'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('首页采用纵向滚动布局，从上到下依次为：顶部问候语（根据时间显示「早上好/下午好/晚上好」）、天气信息卡片、班级统计概览、快捷操作入口、积分排行榜、今日课程表、今日生日学生。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.2 天气信息卡片'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：首页顶部，问候语下方。显示内容包括城市名称、天气图标（晴/阴/雨/雪等）、温度数值、天气描述（如「多云转晴」）。点击卡片可刷新天气数据，加载状态显示旋转动画，错误状态显示默认天气信息。数据来源为调用getWeather云函数，基于用户位置获取实时天气。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.3 班级统计概览'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：天气卡片下方。教师视角显示学生人数（当前班级学生总数）、平均积分（班级学生积分平均值）、考勤率（今日出勤百分比）、待审批（待处理的审批申请数量）。学生/家长视角显示我的积分（当前积分值）、班级排名（积分排名位置）、今日考勤（出勤状态）、待完成任务（值日、活动等）。点击统计卡片跳转对应详情页面，数字动画在页面加载时从0滚动到实际值。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.4 快捷操作入口'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：统计概览下方，横向滑动区域。根据用户角色动态显示不同入口。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('班主任/管理员快捷入口'), font_name='黑体', font_size=12, bold=True)

quick_table = doc.add_table(rows=1, cols=4)
quick_table.style = 'Table Grid'
hdr_cells = quick_table.rows[0].cells
quick_headers = ['图标', '文字', '跳转页面', '功能说明']
for i, header in enumerate(quick_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

quick_items = [
    ('➕', '添加积分', 'score/add/add', '快速登记学生积分'),
    ('📋', '考勤登记', 'attendance/record/record', '记录今日考勤'),
    ('🏠', '宿舍评分', 'dorm/score/add/add', '宿舍卫生检查评分'),
    ('🧹', '值日检查', 'duty/check/check', '检查值日完成情况'),
    ('📢', '发布通知', 'usercenter/notifications', '发送班级通知')
]

for item in quick_items:
    row_cells = quick_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('班干部快捷入口（需授权）'), font_name='黑体', font_size=12, bold=True)

cadre_items = [
    ('➕', '登记积分', 'score/add/add', '登记同学积分变动'),
    ('📋', '考勤登记', 'attendance/record/record', '记录考勤情况'),
    ('🏠', '宿舍评分', 'dorm/score/add/add', '宿舍卫生评分'),
    ('🧹', '值日检查', 'duty/check/check', '值日完成情况检查')
]

for item in cadre_items:
    row_cells = quick_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('学生/家长快捷入口'), font_name='黑体', font_size=12, bold=True)

student_items = [
    ('📊', '我的积分', 'score/score', '查看积分详情'),
    ('🤝', '志愿服务', 'volunteer/volunteer', '报名志愿活动'),
    ('🌳', '心灵树洞', 'treehole/treehole', '匿名倾诉交流'),
    ('🏆', '英雄台', 'hero/hero', '查看荣誉榜单')
]

for item in student_items:
    row_cells = quick_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('交互说明：横向滑动查看更多快捷入口，点击图标跳转对应功能页面，长按图标显示功能说明提示。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.5 积分排行榜'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：快捷入口下方。显示标题「积分排行榜」，前三名显示大头像、姓名、积分值，我的排名显示当前用户/孩子的排名，点击查看更多跳转积分排名完整页面。点击学生头像查看学生详情，下拉刷新更新排行榜数据。排序规则按current_score降序排列，同分处理按bonus_score降序再按学号升序，更新频率为实时更新。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.6 今日课程表'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：排行榜下方。显示标题「今日课程」，课程列表包含时间、科目、教师、教室，课程状态分为未开始/进行中/已结束。点击课程项查看课程详情，可切换班级课程/个人课程视图，显示/隐藏已完成课程。数据来源为课表管理模块，状态判断基于当前时间自动计算，当前进行中的课程蓝色高亮显示。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('3.7 今日生日学生'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：课程表下方。显示标题「今日生日」，学生头像和姓名、年龄信息。点击学生查看学生详情，教师角色显示祝福按钮。'), font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 四、班级模块详解 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('四、班级模块详解'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('4.1 页面结构'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('班级页面采用双栏布局：顶部为搜索栏+筛选条件，中部为班级功能导航（教师角色显示），下部为学生列表。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('4.2 搜索与筛选'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('搜索栏'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：页面顶部。占位文字为「搜索学生姓名或学号」，输入方式为键盘输入支持模糊搜索，搜索按钮为右侧放大镜图标，清除按钮在输入内容后显示点击清空。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('筛选条件（教师角色）'), font_name='黑体', font_size=12, bold=True)

filter_table = doc.add_table(rows=1, cols=3)
filter_table.style = 'Table Grid'
hdr_cells = filter_table.rows[0].cells
filter_headers = ['筛选项', '选项', '说明']
for i, header in enumerate(filter_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

filters = [
    ('班级', '全部班级/具体班级', '按班级筛选学生'),
    ('住宿状态', '全部/住宿生/非住宿', '按住宿状态筛选'),
    ('班干部', '全部/班长/学习委员/...', '按职务筛选')
]

for f in filters:
    row_cells = filter_table.add_row().cells
    for i, text in enumerate(f):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('交互说明：点击筛选标签弹出底部选择面板，选择后自动刷新列表，支持组合筛选条件。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('4.3 班级功能导航'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('位置：搜索栏下方（仅教师角色显示）。功能入口横向排列，5个图标一行，超出可滑动。'), font_name='宋体', font_size=10.5)

nav_table2 = doc.add_table(rows=1, cols=4)
nav_table2.style = 'Table Grid'
hdr_cells = nav_table2.rows[0].cells
nav2_headers = ['图标', '文字', '跳转页面', '功能说明']
for i, header in enumerate(nav2_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

nav2_items = [
    ('📋', '考勤', 'attendance/attendance', '考勤管理首页'),
    ('🧹', '值日', 'duty/duty', '值日管理首页'),
    ('🏠', '宿舍', 'dorm/dorm', '宿舍管理首页'),
    ('⚠️', '处分', 'discipline/record/record', '处分记录管理'),
    ('👥', '分组', 'group/group', '学生分组管理')
]

for item in nav2_items:
    row_cells = nav_table2.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('4.4 学生列表'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('显示内容包括学生头像（默认头像或自定义头像）、姓名、学号、当前积分（带颜色标识）、住宿标识（住宿生显示床位图标）、班干部标识（显示职务标签）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('积分颜色标识'), font_name='黑体', font_size=12, bold=True)

color_table = doc.add_table(rows=1, cols=3)
color_table.style = 'Table Grid'
hdr_cells = color_table.rows[0].cells
color_headers = ['分数范围', '颜色', '等级']
for i, header in enumerate(color_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

color_items = [
    ('≥120', '绿色', '优秀'),
    ('100-119', '蓝色', '良好'),
    ('80-99', '橙色', '一般'),
    ('<80', '红色', '预警')
]

for item in color_items:
    row_cells = color_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('交互说明：点击学生项进入学生详情页面，长按学生项弹出操作菜单（编辑/删除/查看档案），上拉加载分页加载更多学生，下拉刷新刷新学生列表。学生/家长视角仅显示本人/孩子的信息卡片，显示详细个人信息和积分变动。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('4.5 学生详情页面'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：班级页面→点击学生→学生详情。显示内容包括基本信息（头像、姓名、学号、性别、生日）、联系信息（家长电话、家庭住址）、班级信息（班级、职务、住宿状态）、积分信息（当前积分、积分等级、排名）。操作按钮包括编辑信息（仅教师）、查看积分记录、查看考勤记录。点击「编辑信息」进入编辑页面（仅教师），点击「积分记录」查看该学生所有积分变动，点击「考勤记录」查看该学生考勤历史。'), font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 五、积分模块详解 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('五、积分模块详解'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.1 页面结构'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('积分页面根据角色显示不同内容。教师/班干部视角：顶部为班级积分统计，中部为积分排行榜，下部为操作按钮（添加积分、查看记录）。学生/家长视角：顶部为个人积分卡片，中部为积分变动记录，下部为班级排名信息。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.2 积分排行榜（教师视角）'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('显示内容包括排名序号、学生姓名、当前积分、积分等级标签、积分趋势箭头（上升/下降/持平）。排序规则按current_score降序，同分按bonus_score降序再按student_id升序。点击学生查看该学生积分详情，点击表头切换排序方式，搜索框快速查找学生。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.3 添加积分'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：积分页面→点击「添加积分」按钮。操作步骤：选择学生（单选/多选）→选择积分项目（预设项目或自定义）→输入积分值（正数为加分，负数为扣分）→填写原因说明→选择日期（默认当天）→上传佐证图片（可选）→点击「提交」。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('积分项目分类'), font_name='黑体', font_size=12, bold=True)

item_table = doc.add_table(rows=1, cols=3)
item_table.style = 'Table Grid'
hdr_cells = item_table.rows[0].cells
item_headers = ['分类', '示例项目', '默认分值']
for i, header in enumerate(item_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

score_items = [
    ('学习表现', '作业优秀、课堂积极', '+1~+3'),
    ('纪律表现', '迟到、早退、旷课', '-1~-5'),
    ('卫生值日', '值日优秀、值日缺勤', '+1~-2'),
    ('好人好事', '助人为乐、拾金不昧', '+2~+5'),
    ('活动参与', '参加活动、获得奖项', '+1~+10'),
    ('宿舍表现', '卫生优秀、违纪', '+1~-3')
]

for item in score_items:
    row_cells = item_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('审批流程'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('班干部提交需班主任审核，班主任提交直接生效。审核结果分为通过/驳回，通知提交人。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.4 积分记录'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('显示内容包括日期时间、学生姓名、积分变动（+/-数值）、变动原因、操作人、审核状态（待审核/已通过/已驳回）。筛选条件包括日期范围、学生姓名、积分类型（加分/扣分）、审核状态。点击记录查看详情，长按记录弹出操作菜单（编辑/删除，仅管理员），导出按钮导出Excel文件。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.5 积分商城'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/score/mall/mall。功能包括商品展示（积分兑换商品列表）、商品详情（名称、图片、所需积分、库存）、兑换操作（点击兑换扣除积分）、兑换记录（查看历史兑换）。班主任管理功能包括添加商品、编辑商品信息、设置库存、查看兑换台账。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.6 积分申诉'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/score/appeal/appeal。学生/家长可对积分变动提出申诉，填写申诉理由，上传佐证材料，班主任审核申诉，审核结果通知申请人。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('5.7 积分规则配置'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/score/rules/rules。配置内容包括基础分设置（默认100分）、积分项目增删改、积分有效期设置、积分折算比例、规则版本管理。'), font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 六、发现模块详解 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('六、发现模块详解'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('6.1 页面结构'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('发现页面采用分组卡片布局，分为三大板块：学习成长、校园生活、工具服务。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('6.2 学习成长板块'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('成绩管理'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg3/grade/grade。功能为考试成绩记录与分析，操作包括添加成绩、查看成绩趋势、班级成绩对比。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('微聊陪伴（AI聊天）'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg3/aichat/aichat。功能为AI智能对话伙伴，操作为输入文字与AI对话，获取学习建议、心理疏导。家长角色不可见。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('6.3 校园生活板块'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('聚光点（校园活动）'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg3/activity/list/list。功能为校园活动发布与报名，操作包括查看活动列表、点击报名、查看我的活动、发布活动（教师）、审核活动（教师）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('心灵树洞'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg3/treehole/treehole。功能为匿名倾诉交流平台，操作包括发表匿名帖子、选择心情标签、回复他人帖子、点赞支持、举报违规内容。家长角色不可见。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('英雄台'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg3/hero/hero。功能为荣誉殿堂表彰优秀学生，操作包括查看荣誉榜单、颁发荣誉证书（教师）、批量颁发（教师）、点赞鼓励、撤销荣誉（教师）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('校园闲鱼（跳蚤市场）'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg3/flea/flea。功能为闲置物品交易，操作包括发布闲置物品、浏览商品列表、联系卖家、标记售出、举报商品。家长角色不可见。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('6.4 工具服务板块'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('志愿服务'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/volunteer/volunteer。功能为志愿服务记录与管理，操作包括查看志愿活动、报名参加活动、记录服务时长、上传服务证明、审核志愿记录（教师）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('值日管理'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/duty/duty。功能为值日安排与检查，操作包括查看值日表、安排值日（教师）、检查值日完成情况、记录检查结果、查看我的值日（学生）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('纪律管理'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/discipline/record/record。功能为处分记录与撤销，操作包括查看处分记录、添加处分（教师）、申请撤销（学生）、审核撤销申请（教师）、配置处分级别（教师）。'), font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('审批中心'), font_name='黑体', font_size=12, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg1/approval/approval。功能为各类申请审批集中处理，操作包括查看待审批列表、查看已审批列表、审批通过/驳回、查看审批详情。'), font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 七、个人中心详解 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('七、个人中心详解'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('7.1 页面结构'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('个人中心采用分组菜单布局：顶部为用户信息卡片，中部为功能菜单分组，底部为系统设置。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('7.2 用户信息卡片'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('显示内容包括头像（微信头像或自定义）、昵称/姓名、角色标签、班级信息、会员等级标识。点击头像更换头像，点击信息区域编辑个人资料。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('7.3 消息中心'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg5/usercenter/notifications/notifications。功能包括接收系统通知、审批提醒、积分变动提醒、班级通知。'), font_name='宋体', font_size=10.5)

msg_table = doc.add_table(rows=1, cols=3)
msg_table.style = 'Table Grid'
hdr_cells = msg_table.rows[0].cells
msg_headers = ['图标', '类型', '说明']
for i, header in enumerate(msg_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

msg_types = [
    ('📢', '系统通知', '平台公告、更新提醒'),
    ('✅', '审批提醒', '待审批/审批结果通知'),
    ('➕', '积分变动', '积分加减提醒'),
    ('📋', '班级通知', '班主任发布的通知'),
    ('⚠️', '预警提醒', '积分预警、考勤预警')
]

for item in msg_types:
    row_cells = msg_table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

p = doc.add_paragraph()
set_chinese_font(p.add_run('交互说明：点击消息查看详情，左滑消息标记已读/删除，全部已读一键标记所有消息已读。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('7.4 功能菜单'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('菜单分组根据角色动态显示。班主任菜单包括管理功能（发布通知、处分管理、活动审核）、班级设置（班级信息、学期管理、规则配置）、数据管理（数据导入、数据导出、数据迁移）。科任老师菜单包括教学功能（学科积分、学生评价）、个人功能（我的通知、设置）。学生菜单包括我的记录（我的积分、我的考勤、我的值日）、我的申请（我的审批、我的申诉）、成长档案（成绩记录、技能证书、志愿服务）。家长菜单包括孩子信息（积分查看、考勤查看）、家校沟通（班级通知、私信老师）。'), font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('7.5 设置页面'), font_name='黑体', font_size=14, bold=True)

p = doc.add_paragraph()
set_chinese_font(p.add_run('路径：subPkg5/usercenter/settings/settings。设置项包括字体大小（小/中/大）、消息通知（开启/关闭）、积分变动提醒（开启/关闭）、考勤提醒（开启/关闭）、夜间模式（开启/关闭）、清除缓存（按钮）、关于我们（链接）、退出登录（按钮）。'), font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 八、角色专属功能 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('八、角色专属功能'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('8.1 班主任专属功能'), font_name='黑体', font_size=14, bold=True)

features = [
    '班级管理：创建/编辑班级信息、生成/重置班级码、设置班级规则、管理班级成员',
    '学生管理：添加学生（单个/批量导入）、编辑学生信息、删除学生、设置班干部',
    '积分规则配置：设置积分项目、配置积分分值、设置积分有效期、管理规则版本',
    '审批管理：审批积分申请、审批处分撤销、审批活动申请、审批申诉',
    '数据管理：导入学生数据（Excel）、导出成绩单、导出考勤记录、数据迁移',
    '学期管理：创建学期、设置学期时间、激活学期、学期数据归档'
]

for feature in features:
    p = doc.add_paragraph(feature, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('8.2 科任老师专属功能'), font_name='黑体', font_size=14, bold=True)

features2 = [
    '学科积分管理：添加学科积分、查看学科积分统计、学科积分排名',
    '学生评价：写阶段性评语、查看评价历史、评价模板管理',
    '考勤记录：记录课堂考勤、查看考勤统计、设置考勤规则'
]

for feature in features2:
    p = doc.add_paragraph(feature, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('8.3 班干部专属功能（需授权）'), font_name='黑体', font_size=14, bold=True)

features3 = [
    '积分登记：登记同学积分变动、查看登记记录、查看审核状态',
    '考勤管理：记录班级考勤、查看考勤统计、管理请假申请',
    '宿舍评分：宿舍卫生检查、记录评分结果、查看宿舍排名',
    '值日检查：检查值日完成情况、记录检查结果、查看值日统计'
]

for feature in features3:
    p = doc.add_paragraph(feature, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('8.4 学生专属功能'), font_name='黑体', font_size=14, bold=True)

features4 = [
    '个人信息查看：查看个人档案、查看积分详情、查看考勤记录、查看值日安排',
    '活动参与：报名志愿活动、参加校园活动、查看活动记录',
    '申诉申请：积分申诉、处分撤销申请、查看申请状态',
    '社交互动：心灵树洞发帖、英雄台点赞、跳蚤市场交易、AI聊天'
]

for feature in features4:
    p = doc.add_paragraph(feature, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('8.5 家长专属功能'), font_name='黑体', font_size=14, bold=True)

features5 = [
    '孩子信息查看：查看孩子积分、查看积分变动、查看考勤情况、查看值日安排',
    '家校沟通：接收班级通知、私信班主任、查看班级动态',
    '成长关注：查看成绩记录、查看技能证书、查看志愿服务、查看成长档案'
]

for feature in features5:
    p = doc.add_paragraph(feature, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 九、数据流转与业务逻辑 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('九、数据流转与业务逻辑'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('9.1 用户登录流程'), font_name='黑体', font_size=14, bold=True)

flow = [
    '用户打开小程序',
    '检查本地登录状态（openid）',
    '已登录→恢复用户信息→进入首页',
    '未登录→跳转登录页面→微信授权→创建用户档案→选择角色→加入班级→进入首页'
]

for step in flow:
    p = doc.add_paragraph(step, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('9.2 积分变动流程'), font_name='黑体', font_size=14, bold=True)

flow2 = [
    '操作人发起积分变动',
    '填写积分信息（学生、分值、原因）',
    '提交申请',
    '判断操作人角色',
    '班主任/管理员→直接生效→更新学生积分→记录积分日志→通知相关人员',
    '班干部→进入待审核→班主任审核→通过/驳回→更新状态→通知相关人员'
]

for step in flow2:
    p = doc.add_paragraph(step, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('9.3 考勤记录流程'), font_name='黑体', font_size=14, bold=True)

flow3 = [
    '考勤负责人打开考勤页面',
    '选择日期和课程',
    '标记学生出勤状态（出勤/迟到/早退/旷课/请假）',
    '系统自动计算考勤率',
    '异常考勤触发预警',
    '更新学生考勤统计',
    '通知班主任和家长'
]

for step in flow3:
    p = doc.add_paragraph(step, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('9.4 宿舍评分流程'), font_name='黑体', font_size=14, bold=True)

flow4 = [
    '检查人员进入宿舍评分页面',
    '选择宿舍楼栋和房间',
    '按检查项评分（卫生/内务/安全等）',
    '系统自动计算总分',
    '判断是否低于预警线',
    '低于预警线→触发预警通知→通知班主任和宿舍长',
    '积分自动折算→更新宿舍积分账户→更新学生个人积分'
]

for step in flow4:
    p = doc.add_paragraph(step, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('9.5 处分管理流程'), font_name='黑体', font_size=14, bold=True)

flow5 = [
    '班主任记录处分',
    '填写处分信息（学生、级别、原因）',
    '系统自动扣除对应积分',
    '记录处分档案',
    '通知学生和家长',
    '学生申请撤销（达到期限）',
    '班主任审核撤销申请',
    '通过→撤销处分→恢复积分→记录撤销档案',
    '驳回→维持处分→通知学生'
]

for step in flow5:
    p = doc.add_paragraph(step, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('9.6 审批工作流程'), font_name='黑体', font_size=14, bold=True)

flow6 = [
    '申请人提交申请',
    '系统判断申请类型',
    '确定审批人',
    '通知审批人',
    '审批人审核',
    '通过→执行对应操作→通知申请人',
    '驳回→记录驳回原因→通知申请人'
]

for step in flow6:
    p = doc.add_paragraph(step, style='List Number')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 十、常见问题处理 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('十、常见问题处理'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('10.1 登录问题'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题1：无法登录'), font_name='黑体', font_size=12, bold=True)

solutions = [
    '检查网络连接',
    '确认微信版本支持小程序',
    '清除小程序缓存后重试',
    '联系管理员检查账号状态'
]

for sol in solutions:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题2：角色选择错误'), font_name='黑体', font_size=12, bold=True)

solutions2 = [
    '退出当前账号',
    '重新登录',
    '选择正确角色',
    '如已绑定班级，需联系班主任调整角色'
]

for sol in solutions2:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('10.2 班级加入问题'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题1：班级码无效'), font_name='黑体', font_size=12, bold=True)

solutions3 = [
    '确认班级码输入正确（注意大小写）',
    '检查班级码是否已过期',
    '联系班主任获取最新班级码',
    '确认班级状态为「正常」'
]

for sol in solutions3:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题2：加入班级后看不到数据'), font_name='黑体', font_size=12, bold=True)

solutions4 = [
    '检查网络连接',
    '下拉刷新页面',
    '确认角色权限',
    '联系班主任确认是否已审核通过'
]

for sol in solutions4:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('10.3 积分问题'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题1：积分显示不正确'), font_name='黑体', font_size=12, bold=True)

solutions5 = [
    '下拉刷新积分页面',
    '检查积分记录是否有遗漏',
    '确认积分规则配置',
    '联系班主任核实'
]

for sol in solutions5:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题2：积分变动未生效'), font_name='黑体', font_size=12, bold=True)

solutions6 = [
    '检查审批状态（班干部提交需审核）',
    '确认网络连接正常',
    '查看操作记录',
    '重新提交或联系班主任'
]

for sol in solutions6:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('10.4 功能使用问题'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题1：找不到某个功能'), font_name='黑体', font_size=12, bold=True)

solutions7 = [
    '检查当前角色是否有该功能权限',
    '确认功能是否在当前版本中开放',
    '查看「发现」页面是否有入口',
    '联系管理员确认功能状态'
]

for sol in solutions7:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题2：页面加载缓慢'), font_name='黑体', font_size=12, bold=True)

solutions8 = [
    '检查网络连接',
    '清除小程序缓存',
    '关闭其他应用释放内存',
    '稍后重试'
]

for sol in solutions8:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('10.5 数据问题'), font_name='黑体', font_size=14, bold=True)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题1：数据不同步'), font_name='黑体', font_size=12, bold=True)

solutions9 = [
    '下拉刷新页面',
    '退出重新进入小程序',
    '检查网络连接',
    '联系管理员检查服务器状态'
]

for sol in solutions9:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h3 = doc.add_heading('', level=3)
set_chinese_font(h3.add_run('问题2：数据丢失'), font_name='黑体', font_size=12, bold=True)

solutions10 = [
    '检查是否切换了班级',
    '确认登录账号正确',
    '查看操作日志',
    '联系管理员恢复数据'
]

for sol in solutions10:
    p = doc.add_paragraph(sol, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

doc.add_page_break()

# ========== 十一、界面元素速查表 ==========
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('十一、界面元素速查表'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('11.1 图标含义'), font_name='黑体', font_size=14, bold=True)

icon_table = doc.add_table(rows=1, cols=3)
icon_table.style = 'Table Grid'
hdr_cells = icon_table.rows[0].cells
icon_headers = ['图标', '含义', '使用场景']
for i, header in enumerate(icon_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

icons = [
    ('🏠', '首页', '底部导航'),
    ('👥', '班级', '底部导航'),
    ('📊', '积分', '底部导航'),
    ('🔍', '发现', '底部导航'),
    ('👤', '我的', '底部导航'),
    ('➕', '添加', '添加积分、添加学生'),
    ('✏️', '编辑', '编辑信息'),
    ('🗑️', '删除', '删除记录'),
    ('✅', '确认/通过', '审批通过'),
    ('❌', '取消/驳回', '审批驳回'),
    ('📋', '记录', '查看记录'),
    ('📢', '通知', '发布通知'),
    ('⚠️', '警告', '预警提醒'),
    ('🏆', '荣誉', '英雄台'),
    ('🌳', '树洞', '心灵树洞'),
    ('🤝', '志愿', '志愿服务'),
    ('🧹', '值日', '值日管理'),
    ('🏠', '宿舍', '宿舍管理'),
    ('💬', '聊天', 'AI聊天'),
    ('📅', '日期', '日期选择'),
    ('📷', '图片', '上传图片'),
    ('🔔', '提醒', '消息提醒'),
    ('🔄', '刷新', '刷新数据'),
    ('📤', '导出', '导出数据'),
    ('📥', '导入', '导入数据')
]

for icon in icons:
    row_cells = icon_table.add_row().cells
    for i, text in enumerate(icon):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('11.2 颜色含义'), font_name='黑体', font_size=14, bold=True)

color_table2 = doc.add_table(rows=1, cols=3)
color_table2.style = 'Table Grid'
hdr_cells = color_table2.rows[0].cells
color2_headers = ['颜色', '含义', '使用场景']
for i, header in enumerate(color2_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

colors = [
    ('蓝色 #1890ff', '主色调', '按钮、导航、链接'),
    ('紫色 #667eea', '强调色', '选中状态、高亮'),
    ('绿色 #52c41a', '成功', '成功提示、通过状态'),
    ('红色 #ff4d4f', '危险/警告', '删除按钮、扣分、预警'),
    ('橙色 #faad14', '提醒', '待审核、一般预警'),
    ('灰色 #999999', '禁用/次要', '未选中、次要文字')
]

for color in colors:
    row_cells = color_table2.add_row().cells
    for i, text in enumerate(color):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('11.3 状态标签'), font_name='黑体', font_size=14, bold=True)

status_table = doc.add_table(rows=1, cols=3)
status_table.style = 'Table Grid'
hdr_cells = status_table.rows[0].cells
status_headers = ['标签', '颜色', '含义']
for i, header in enumerate(status_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

statuses = [
    ('已通过', '绿色', '审批通过'),
    ('待审核', '橙色', '等待审核'),
    ('已驳回', '红色', '审批未通过'),
    ('已撤销', '灰色', '已撤销/作废'),
    ('进行中', '蓝色', '正在进行'),
    ('已完成', '绿色', '已经完成'),
    ('已过期', '灰色', '超过有效期')
]

for status in statuses:
    row_cells = status_table.add_row().cells
    for i, text in enumerate(status):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

# 附录
h1 = doc.add_heading('', level=1)
set_chinese_font(h1.add_run('附录'), font_name='黑体', font_size=16, bold=True)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('A. 版本更新记录'), font_name='黑体', font_size=14, bold=True)

version_table = doc.add_table(rows=1, cols=3)
version_table.style = 'Table Grid'
hdr_cells = version_table.rows[0].cells
version_headers = ['版本', '日期', '更新内容']
for i, header in enumerate(version_headers):
    hdr_cells[i].text = header
    for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
            set_chinese_font(run, font_name='黑体', font_size=10.5, bold=True)

versions = [
    ('v1.0', '2025-06', '体验版上线，包含核心功能模块')
]

for v in versions:
    row_cells = version_table.add_row().cells
    for i, text in enumerate(v):
        row_cells[i].text = text
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                set_chinese_font(run, font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('B. 技术支持'), font_name='黑体', font_size=14, bold=True)

support_items = [
    '技术问题反馈：请在微信群中@管理员',
    '功能建议：请通过「我的-设置-意见反馈」提交',
    '紧急问题：请联系班主任或系统管理员'
]

for item in support_items:
    p = doc.add_paragraph(item, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

h2 = doc.add_heading('', level=2)
set_chinese_font(h2.add_run('C. 法律声明'), font_name='黑体', font_size=14, bold=True)

legal_items = [
    '本系统所有数据归班级管理方所有',
    '用户需遵守相关法律法规和平台规则',
    '禁止利用系统进行违法违规活动'
]

for item in legal_items:
    p = doc.add_paragraph(item, style='List Bullet')
    set_chinese_font(p.runs[0], font_name='宋体', font_size=10.5)

# 文档结束
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_chinese_font(p.add_run('文档结束\n\n本手册基于班分通小程序体验版 v1.0 编写，正式版发布后可能会有功能调整，请以实际版本为准。'), font_name='宋体', font_size=10.5)

# 保存文档
doc.save('班分通小程序详细操作手册.docx')
print('Word文档已生成成功！')